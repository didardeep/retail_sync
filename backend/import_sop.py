"""
Imports the SOP audit tools (Cash, FMCG) from docs/SOP_Audit_Checklists.xlsx
into the sop_* tables by publishing them through app/sop_versions.py, so the
import and the tool editor create template versions the same way.

A tool with no version yet gets version 1. Otherwise the sheet is matched to
the current version (sections by code, questions by position within the
section) so stable keys and manager-set proof flags (requires_comment /
requires_photo) carry over, and a new version is published only if the sheet
differs from the current one. Running it twice therefore publishes nothing
the second time.

    python import_sop.py
    python import_sop.py path/to/other.xlsx
"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.db import SessionLocal, init_db
from app.models import SopTemplate
from app.sop_versions import publish_version, tree_of

DEFAULT_FILE = Path(__file__).parent.parent / "docs" / "SOP_Audit_Checklists.xlsx"

# sheet title prefix -> (template code, how the sheet defines the Minimum level)
TOOLS = {
    "cash": ("CASH", "zero"),
    "fmcg": ("FMCG", "third"),
}

SECTION_RE = re.compile(r"^Section\s+([A-Z]):\s*(.*?)\s*(?:\|.*)?$", re.S)


def _text(value):
    """Cell -> stripped string, or None for empty / 'NA'."""
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.upper() == "NA":
        return None
    return s


def _tool_for(sheet_title):
    key = sheet_title.strip().split()[0].lower()
    return TOOLS.get(key)


def parse_sheet(ws):
    """Return a list of sections: {code, name, criteria:[{...}]}."""
    sections = []
    current = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        first, marks = row[0], row[1]
        if first is None:
            continue
        title = str(first).strip()
        m = SECTION_RE.match(title)
        if marks is None and m:
            current = {"code": m.group(1), "name": m.group(2), "criteria": []}
            sections.append(current)
            continue
        if current is None or marks is None:
            continue
        current["criteria"].append({
            "title": title,
            "marks": float(marks),
            "max_text": _text(row[2]),
            "avg_text": _text(row[3]),
            "min_text": _text(row[4]),
            "default_na": "(N/A)" in title,
        })
    return sections


def _tree_from_sheet(sections, code, name, rule, current):
    """The sheet as a tree. Where the current version has the same section code
    (and a question at the same position) its stable key and proof flags are
    carried over, so the question stays "the same question" across versions."""
    old = {s["code"]: s for s in tree_of(current)["sections"]} if current else {}
    out = []
    for sec in sections:
        before = old.get(sec["code"])
        criteria = []
        for idx, crit in enumerate(sec["criteria"]):
            prev = before["criteria"][idx] if before and idx < len(before["criteria"]) else None
            criteria.append({
                **crit,
                "key": prev["key"] if prev else None,
                "requires_comment": prev["requires_comment"] if prev else False,
                "requires_photo": prev["requires_photo"] if prev else False,
            })
        out.append({"key": before["key"] if before else None,
                    "code": sec["code"], "name": sec["name"], "criteria": criteria})
    return {"code": code, "name": name, "min_rule": rule, "sections": out}


def upsert_template(db, ws, code, rule, user=None):
    """Publish the sheet as the tool's next version unless nothing changed.
    Returns (template, section count, criterion count, published) where
    `published` is False when the sheet matched the current version."""
    sections = parse_sheet(ws)
    current = db.query(SopTemplate).filter_by(code=code, is_current=True).first()
    tree = _tree_from_sheet(sections, code, ws.title.strip(), rule, current)
    result = publish_version(
        db, code, tree, user, current.version if current else None,
        "Imported from the audit checklist spreadsheet")
    tpl = result.template
    if not tpl.is_active:
        tpl.is_active = True
        db.commit()
    n_criteria = sum(len(s["criteria"]) for s in sections)
    return tpl, len(sections), n_criteria, not result.unchanged


def import_sop(path=DEFAULT_FILE):
    wb = load_workbook(str(path), data_only=True)
    init_db()
    db = SessionLocal()
    try:
        for ws in wb.worksheets:
            tool = _tool_for(ws.title)
            if tool is None:
                print(f"  skipped sheet (unknown tool): {ws.title!r}")
                continue
            tpl, n_sec, n_crit, published = upsert_template(db, ws, *tool)
            note = f"published v{tpl.version}" if published else "unchanged"
            print(f"  {tpl.code}: {n_sec} sections, {n_crit} criteria, "
                  f"{tpl.total_marks:g} marks ({note})")
    finally:
        db.close()
        wb.close()


if __name__ == "__main__":
    print("Importing SOP audit tools...")
    import_sop(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FILE)
    print("Done.")
