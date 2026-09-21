"""
Import Excel files from uploads/ folder into the database tables.

    python import_excel.py            # import all
    python import_excel.py --reset    # drop+recreate tables first, then seed + import
"""
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.db import Base, SessionLocal, engine, init_db
from app.models import (
    CashReconciliation, CashDepositPickup, DataImport,
    ExpiredInventory, StoreScore,
)

UPLOADS = Path(__file__).parent / "uploads"


# ---------------------------------------------------------------------------
# JSON safety — Excel date cells are python datetime, not JSON-serialisable
# ---------------------------------------------------------------------------
def _json_safe(val):
    """Coerce a single value to a JSON-serialisable form."""
    if val is None:
        return None
    if isinstance(val, dt.datetime):
        return val.isoformat()
    if isinstance(val, dt.date):
        return val.isoformat()
    if isinstance(val, (int, float, str, bool)):
        return val
    return str(val)


def json_safe_row(row: dict) -> dict:
    """Make every value in a row dict JSON-serialisable (dates -> isoformat)."""
    return {k: _json_safe(v) for k, v in row.items()}


def _parse_date(val):
    if val is None:
        return None
    if isinstance(val, dt.date):
        return val
    if isinstance(val, dt.datetime):
        return val.date()
    try:
        return dt.date.fromisoformat(str(val).strip())
    except (ValueError, TypeError):
        return None


def _float(val):
    if val is None:
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _int(val):
    if val is None:
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return None


def _read_sheet(path):
    """Read an Excel file and return (headers, rows) where rows are list of dicts."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []
    for row in rows[1:]:
        d = {}
        for i, val in enumerate(row):
            if i < len(headers):
                d[headers[i]] = val
        data.append(d)
    wb.close()
    return headers, data


def _resolve_store_id(session, row):
    """Resolve 'Store ID' to a stores.id, handling ST0xx and numeric store_code."""
    from app.models import Store
    raw = str(row.get("Store ID", "")).strip()
    if not raw:
        return raw
    if raw.upper().startswith("ST"):
        return raw
    store = session.query(Store).filter(Store.store_code == raw).first()
    if store:
        return store.id
    try:
        return f"ST{int(raw):03d}"
    except (ValueError, TypeError):
        return raw


def _normalise_status(val):
    """Canonical status: 'Operational' and 'Operating' both -> 'Operating'."""
    s = str(val or "").strip()
    if s.lower() == "operational":
        return "Operating"
    return s


def import_cash_reconciliation(session):
    path = UPLOADS / "cash_reconciliation.xlsx"
    if not path.exists():
        print(f"  Skipped: {path.name} not found")
        return 0

    headers, data = _read_sheet(path)
    di = DataImport(
        file_name=path.name,
        data_section="cash_reconciliation",
        column_headers=headers,
    )
    session.add(di)
    session.flush()

    for idx, row in enumerate(data, 1):
        sid = _resolve_store_id(session, row)
        session.add(CashReconciliation(
            import_id=di.id,
            store_id=sid,
            source_row_number=idx,
            cash_at_tills=_float(row.get("Cash at Tills")),
            cash_in_safe=_float(row.get("Cash in Safe")),
            cash_other_locations=_float(row.get("Cash Other Locations") or row.get("Any other place")),
            physical_cash_total=_float(row.get("Physical Cash Total")),
            cash_sales_as_per_report=_float(row.get("Cash Sales as per Report") or row.get("Cash Sales as per sales Report")),
            float_or_imprest_amount=_float(row.get("Float or Imprest Amount") or row.get("Float and Imprest allocated to the store as per Master")),
            book_cash_total=_float(row.get("Book Cash Total")),
            difference=_float(row.get("Difference") or row.get("Difference (A-B)")),
            remarks=str(row.get("Remarks", "") or ""),
            raw_data=json_safe_row(row),
        ))

    print(f"  Imported {len(data)} cash reconciliation records")
    return len(data)


def import_cash_deposit_pickups(session):
    path = UPLOADS / "cash_deposit_pickups.xlsx"
    if not path.exists():
        print(f"  Skipped: {path.name} not found")
        return 0

    headers, data = _read_sheet(path)
    di = DataImport(
        file_name=path.name,
        data_section="cash_deposit_pickups",
        column_headers=headers,
    )
    session.add(di)
    session.flush()

    for idx, row in enumerate(data, 1):
        sid = _resolve_store_id(session, row)
        session.add(CashDepositPickup(
            import_id=di.id,
            store_id=sid,
            source_row_number=idx,
            sales_date=_parse_date(row.get("Sales Date") or row.get("Sales Date (A)")),
            cash_sales=_float(row.get("Cash Sales")),
            cash_deposited=_float(row.get("Cash Deposited") or row.get("Cash Deposited (c)")),
            difference=_float(row.get("Difference") or row.get("Difference (B-C)")),
            cms_pickup_date=_parse_date(row.get("CMS Pickup Date") or row.get("CMS Pick Up date (D)")),
            handover_delay_days=_int(row.get("Handover Delay Days") or row.get("Delay in cash handover (A-D)")),
            remarks=str(row.get("Remarks", "") or ""),
            raw_data=json_safe_row(row),
        ))

    print(f"  Imported {len(data)} cash deposit pickup records")
    return len(data)


def import_expired_inventory(session):
    path = UPLOADS / "expired_inventory.xlsx"
    if not path.exists():
        print(f"  Skipped: {path.name} not found")
        return 0

    headers, data = _read_sheet(path)
    di = DataImport(
        file_name=path.name,
        data_section="expired_inventory",
        column_headers=headers,
    )
    session.add(di)
    session.flush()

    for idx, row in enumerate(data, 1):
        sid = _resolve_store_id(session, row)
        session.add(ExpiredInventory(
            import_id=di.id,
            store_id=sid,
            source_row_number=idx,
            article_code=str(row.get("Article Code", "") or ""),
            article_description=str(row.get("Article Description", "") or ""),
            expiry_date=_parse_date(row.get("Expiry Date")),
            review_date=_parse_date(row.get("Review Date")),
            quantity=_float(row.get("Quantity")),
            mrp=_float(row.get("MRP")),
            raw_data=json_safe_row(row),
        ))

    print(f"  Imported {len(data)} expired inventory records")
    return len(data)


def import_store_scores(session):
    path = UPLOADS / "store_scores.xlsx"
    if not path.exists():
        print(f"  Skipped: {path.name} not found")
        return 0

    headers, data = _read_sheet(path)
    di = DataImport(
        file_name=path.name,
        data_section="store_scores",
        column_headers=headers,
    )
    session.add(di)
    session.flush()

    for idx, row in enumerate(data, 1):
        sid = _resolve_store_id(session, row)
        session.add(StoreScore(
            import_id=di.id,
            store_id=sid,
            source_row_number=idx,
            store_code=str(row.get("Store Code") or row.get("Store code") or ""),
            status=_normalise_status(row.get("Status")),
            q1=_float(row.get("Q1")),
            q2=_float(row.get("Q2")),
            q3=_float(row.get("Q3")),
            q4=_float(row.get("Q4")),
            raw_data=json_safe_row(row),
        ))

    print(f"  Imported {len(data)} store score records")
    return len(data)


def run_import():
    session = SessionLocal()
    try:
        print("Importing Excel files from uploads/...")
        total = 0
        total += import_cash_reconciliation(session)
        total += import_cash_deposit_pickups(session)
        total += import_expired_inventory(session)
        total += import_store_scores(session)
        session.commit()
        print(f"\nDone! Imported {total} total records.")
    except Exception as e:
        session.rollback()
        print(f"Error: {e}")
        raise
    finally:
        session.close()


if __name__ == "__main__":
    if "--reset" in sys.argv:
        Base.metadata.drop_all(bind=engine)
        print("Dropped all tables.")
    init_db()

    # If --reset, seed first then import
    if "--reset" in sys.argv:
        print("Running seed first...")
        from seed import seed
        seed()
        print()

    run_import()
