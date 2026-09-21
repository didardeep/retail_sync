"""
Import master Excel files into the database:

  - Store_Master_dummy.xlsx  -> stores, store_scores, observations
  - Store_Audit_Checklist_v1_dummy.xlsx -> questions, cash_recon, cash_deposit, expired_inventory

    python import_master.py
"""
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.auth import hash_password
from app.db import Base, SessionLocal, engine, init_db
from app.models import (
    ROLE_AUDIT_MANAGER, ROLE_AUDITOR, ROLE_STORE_MANAGER,
    Audit, AuditResponse, AuditorAvailability, CashReconciliation,
    CashDepositPickup, Checklist, ChecklistItem, DataImport,
    ExpiredInventory, Issue, Observation, Question, Store, StoreScore, User,
)
from import_excel import (
    _parse_date, _float, _int, _normalise_status,
    json_safe_row, _read_sheet,
)

UPLOADS = Path(__file__).parent / "uploads"
MASTER_FILE = UPLOADS / "Store_Master_dummy.xlsx"
CHECKLIST_FILE = UPLOADS / "Store_Audit_Checklist_v1_dummy.xlsx"
DEFAULT_PASSWORD = "password123"


def _read_named_sheet(path, sheet_name):
    """Read a specific named sheet from an Excel workbook.
    Skips leading blank rows and uses the first row with real data as headers.
    Also skips trailing all-None data rows.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return [], []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []

    # Skip leading blank rows to find the real header row
    hdr_idx = 0
    for i, row in enumerate(rows):
        non_empty = [v for v in row if v is not None and str(v).strip()]
        if len(non_empty) >= 2:  # at least 2 non-empty cells = likely header
            hdr_idx = i
            break

    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[hdr_idx])]
    data = []
    for row in rows[hdr_idx + 1:]:
        # skip all-empty rows
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        d = {}
        for i, val in enumerate(row):
            if i < len(headers):
                d[headers[i]] = val
        data.append(d)
    return headers, data


def _resolve_sid(row, code_to_id, store_map):
    """Try to find a valid store_id from a data row's Store ID / Store code columns."""
    for key in ("Store ID", "Store Id", "Store code", "Store Code"):
        raw = str(row.get(key, "") or "").strip()
        if not raw:
            continue
        # Direct match
        if raw in store_map:
            return raw
        # Via code_to_id lookup
        sid = code_to_id.get(raw)
        if sid and sid in store_map:
            return sid
    return None


def _json_safe_val(val):
    if val is None:
        return None
    if isinstance(val, dt.datetime):
        return val.isoformat()
    if isinstance(val, dt.date):
        return val.isoformat()
    if isinstance(val, (int, float, str, bool)):
        return val
    return str(val)


def _json_safe(row):
    return {k: _json_safe_val(v) for k, v in row.items()}


def seed_from_excel():
    """Load everything from Excel files instead of seed_data.json."""
    s = SessionLocal()

    # ====================================================================
    # 1. Users (create a fixed set of users)
    # ====================================================================
    admin = User(
        name="Audit Manager", email="am@retail-chain.com",
        role=ROLE_AUDIT_MANAGER, designation="Audit Manager",
        password_hash=hash_password(DEFAULT_PASSWORD),
    )
    s.add(admin)
    s.flush()

    auditor_names = ["J. Patel", "A. Sharma", "P. Rao", "N. Singh", "R. Das"]
    auditors = {}
    for name in auditor_names:
        slug = ".".join(p for p in name.lower().replace(".", " ").split() if p)
        u = User(
            name=name, email=f"aud.{slug}@retail-chain.com",
            role=ROLE_AUDITOR, designation="Field Auditor",
            password_hash=hash_password(DEFAULT_PASSWORD),
        )
        s.add(u)
        s.flush()
        auditors[name] = u

    managers = {}  # name -> User

    def get_or_create_manager(name):
        if not name:
            return None
        if name in managers:
            return managers[name]
        slug = ".".join(p for p in name.lower().replace(".", " ").split() if p)
        u = User(
            name=name, email=f"sm.{slug}@retail-chain.com",
            role=ROLE_STORE_MANAGER, designation="Store Manager",
            password_hash=hash_password(DEFAULT_PASSWORD),
        )
        s.add(u)
        s.flush()
        managers[name] = u
        return u

    # ====================================================================
    # 2. Stores from Store_Master_dummy.xlsx "Store Details"
    # ====================================================================
    print("\n--- Loading stores from Store_Master_dummy.xlsx ---")
    headers, store_rows = _read_named_sheet(MASTER_FILE, "Store Details")
    print(f"  Found {len(store_rows)} store rows, headers: {headers}")

    store_map = {}  # store_id -> Store
    code_to_id = {}  # store_code -> store_id

    for row in store_rows:
        sid = str(row.get("Store ID", "")).strip()
        code = str(row.get("Store code", "")).strip()
        name = str(row.get("Name", "")).strip()
        mgr_name = str(row.get("Operation Manager") or row.get("Deputy Manager") or "").strip()

        mgr = get_or_create_manager(mgr_name) if mgr_name else None

        status = _normalise_status(row.get("Status"))

        store = Store(
            id=sid,
            store_code=code,
            name=name,
            state=str(row.get("State", "") or ""),
            city=str(row.get("City", "") or ""),
            region=str(row.get("Region", "") or ""),
            district=str(row.get("District", "") or ""),
            postal_code=str(row.get("Postal Address", "") or ""),
            deputy_manager=str(row.get("Deputy Manager", "") or ""),
            operations_manager=str(row.get("Operation Manager", "") or ""),
            store_category=str(row.get("StoreCatg_main", "") or ""),
            status=status,
            manager_id=mgr.id if mgr else None,
            meta={},
        )
        s.add(store)
        store_map[sid] = store
        code_to_id[code] = sid

    s.flush()
    print(f"  Loaded {len(store_map)} stores")

    # ====================================================================
    # 3. Store Scores from Store_Master_dummy.xlsx "Store Scores"
    # ====================================================================
    print("\n--- Loading store scores ---")
    headers, score_rows = _read_named_sheet(MASTER_FILE, "Store Scores")
    di_scores = DataImport(
        file_name="Store_Master_dummy.xlsx",
        data_section="store_scores",
        column_headers=headers,
    )
    s.add(di_scores)
    s.flush()

    loaded_scores = 0
    for idx, row in enumerate(score_rows, 1):
        code = str(row.get("Store code") or row.get("Store Code") or "").strip()
        if not code:
            continue  # skip rows without a store code
        sid = code_to_id.get(code, code)
        # Only insert if store_id exists in stores table
        if sid not in store_map:
            print(f"    Skipping score row {idx}: store code '{code}' not found in stores")
            continue
        s.add(StoreScore(
            import_id=di_scores.id,
            store_id=sid,
            source_row_number=idx,
            store_code=code,
            status=_normalise_status(row.get("Status")),
            q1=_float(row.get("Q1")),
            q2=_float(row.get("Q2")),
            q3=_float(row.get("Q3")),
            q4=_float(row.get("Q4")),
            raw_data=_json_safe(row),
        ))
        loaded_scores += 1
    s.flush()
    print(f"  Loaded {loaded_scores} store score rows")

    # ====================================================================
    # 4. Observations from Store_Master_dummy.xlsx "Observation"
    # ====================================================================
    print("\n--- Loading observations ---")
    headers, obs_rows = _read_named_sheet(MASTER_FILE, "Observation")
    first_store = list(store_map.keys())[0] if store_map else None

    loaded_obs = 0
    for row in obs_rows:
        obs_text = str(row.get("Observation", "") or "").strip()
        if not obs_text:
            continue  # skip empty observation rows
        # Try to resolve store from row if a Store ID/Code column exists
        obs_store = None
        for k in ("Store ID", "Store Id", "Store code", "Store Code"):
            v = str(row.get(k, "") or "").strip()
            if v:
                obs_store = code_to_id.get(v, store_map.get(v, {}) and v)
                if obs_store in store_map:
                    break
                obs_store = None
        if not obs_store:
            obs_store = first_store  # fallback to first store

        s.add(Observation(
            store_id=obs_store,
            sr_no=_int(row.get("Sr No") or row.get("Sr. No") or row.get("S.No")),
            observation=obs_text,
            risk=str(row.get("Risk", "") or ""),
            action_plan=str(row.get("Action Plan", "") or ""),
            person_responsible=str(row.get("Person Responsible", "") or ""),
            target=str(row.get("Target", "") or ""),
            status=str(row.get("Status", "") or ""),
        ))
        loaded_obs += 1
    s.flush()
    print(f"  Loaded {loaded_obs} observations")

    # ====================================================================
    # 5. Questions from Store_Audit_Checklist_v1_dummy.xlsx "All Process Checklist"
    # ====================================================================
    print("\n--- Loading questions from checklist ---")
    headers, q_rows = _read_named_sheet(CHECKLIST_FILE, "All Process Checklist")

    questions = {}
    for idx, row in enumerate(q_rows, 1):
        code = f"Q{idx:03d}"
        text = str(row.get("Question", "")).strip()
        if not text:
            continue
        q = Question(
            code=code, version=1, is_current=True,
            text=text,
            process=str(row.get("Process", "") or ""),
            sub_process=str(row.get("Sub-Process", "") or ""),
            audit_type=str(row.get("Type", "") or ""),
            weight=_int(row.get("Weights")) or 1,
            is_critical=(_int(row.get("Weights")) or 0) >= 5,
            active=True,
            photo_video_enablement=str(row.get("Photo/Video", "No") or "No"),
            data_analyst_enabled=str(row.get("Data Analyst Enabled", "N") or "N"),
            data_analyst_reference=str(row.get("Data Analyst_Reference", "") or "") or None,
            annex=str(row.get("Annexure", "") or "") or None,
            approval_status="APPROVED",
            approved_by_id=admin.id,
            meta={},
        )
        s.add(q)
        s.flush()
        questions[code] = q

    print(f"  Loaded {len(questions)} questions")

    # ====================================================================
    # 6. Checklists — one per process + a full checklist
    # ====================================================================
    processes = sorted({q.process for q in questions.values() if q.process})
    checklists = {}
    for proc in processes:
        cl = Checklist(name=f"{proc} Checklist", version=1,
                       conditions={"process": [proc]})
        s.add(cl)
        s.flush()
        items = [q for q in questions.values() if q.process == proc]
        for i, q in enumerate(items):
            s.add(ChecklistItem(checklist_id=cl.id, question_id=q.id, sort_order=i))
        checklists[proc] = cl

    full = Checklist(name="Full Store Audit", version=1, conditions={})
    s.add(full)
    s.flush()
    for i, q in enumerate(questions.values()):
        s.add(ChecklistItem(checklist_id=full.id, question_id=q.id, sort_order=i))
    s.flush()

    # ====================================================================
    # 7. Cash Reconciliation from checklist sheet "1"
    #    Layout can be: (a) tabular with Store ID column, or
    #    (b) vertical form (Particulars / Amount) for a single store visit
    # ====================================================================
    print("\n--- Loading cash reconciliation from checklist sheet '1' ---")
    di_cr = DataImport(
        file_name="Store_Audit_Checklist_v1_dummy.xlsx",
        data_section="cash_reconciliation",
        column_headers=[],
    )
    s.add(di_cr)
    s.flush()

    loaded_cr = 0
    first_store_id = list(store_map.keys())[0] if store_map else None

    # Read raw rows to detect layout
    wb_cr = load_workbook(str(CHECKLIST_FILE), read_only=True, data_only=True)
    if "1" in wb_cr.sheetnames:
        ws_cr = wb_cr["1"]
        raw_cr = list(ws_cr.iter_rows(values_only=True))
        wb_cr.close()

        # Detect form layout: look for "Particulars" or vertical key-value pairs
        is_form = any(
            "Particulars" in str(r[1] or "") or "Cash at Tills" in str(r[1] or "")
            for r in raw_cr[:5] if len(r) > 1
        )

        if is_form:
            # Parse vertical form into one CashReconciliation record
            vals = {}
            for r in raw_cr:
                label = str(r[1] or "").strip() if len(r) > 1 else ""
                amount = r[2] if len(r) > 2 else None
                amount_total = r[3] if len(r) > 3 else None
                if label:
                    vals[label] = _float(amount)
                if amount_total is not None:
                    # Check if this is a section total or difference
                    section_marker = str(r[0] or "").strip() if r[0] else ""
                    if section_marker == "C" or "Difference" in label:
                        vals["Difference"] = _float(amount_total)
                    elif "Book" in label or section_marker == "B":
                        vals["Book Cash Total"] = _float(amount_total)
                    elif section_marker == "A" or not section_marker:
                        vals["Physical Cash Total"] = _float(amount_total)

            s.add(CashReconciliation(
                import_id=di_cr.id,
                store_id=first_store_id,
                source_row_number=1,
                cash_at_tills=vals.get("Cash at Tills"),
                cash_in_safe=vals.get("Cash in Safe"),
                cash_other_locations=vals.get("Any other place"),
                physical_cash_total=vals.get("Physical Cash Total"),
                cash_sales_as_per_report=vals.get("Cash Sales as per sales Report") or vals.get("Cash Sales as per Report"),
                float_or_imprest_amount=vals.get("Float and Imprest allocated to the store as per Master"),
                book_cash_total=vals.get("Book Cash Total"),
                difference=vals.get("Difference"),
                remarks="",
                raw_data={k: v for k, v in vals.items() if v is not None},
            ))
            loaded_cr = 1
        else:
            # Tabular layout
            headers, cr_rows = _read_named_sheet(CHECKLIST_FILE, "1")
            di_cr.column_headers = headers
            for idx, row in enumerate(cr_rows, 1):
                sid = _resolve_sid(row, code_to_id, store_map) or first_store_id
                if not sid:
                    continue
                s.add(CashReconciliation(
                    import_id=di_cr.id,
                    store_id=sid,
                    source_row_number=idx,
                    cash_at_tills=_float(row.get("Cash at Tills")),
                    cash_in_safe=_float(row.get("Cash in Safe")),
                    cash_other_locations=_float(row.get("Any other place") or row.get("Cash Other Locations")),
                    physical_cash_total=_float(row.get("Physical Cash Total")),
                    cash_sales_as_per_report=_float(row.get("Cash Sales as per sales Report") or row.get("Cash Sales as per Report")),
                    float_or_imprest_amount=_float(row.get("Float and Imprest allocated to the store as per Master") or row.get("Float or Imprest Amount")),
                    book_cash_total=_float(row.get("Book Cash Total")),
                    difference=_float(row.get("Difference (A-B)") or row.get("Difference")),
                    remarks=str(row.get("Remarks", "") or ""),
                    raw_data=_json_safe(row),
                ))
                loaded_cr += 1
    else:
        wb_cr.close()
    print(f"  Loaded {loaded_cr} cash reconciliation rows")

    # ====================================================================
    # 8. Cash Deposit Pickups from checklist sheet "2"
    # ====================================================================
    print("\n--- Loading cash deposit pickups from checklist sheet '2' ---")
    headers, cd_rows = _read_named_sheet(CHECKLIST_FILE, "2")
    di_cd = DataImport(
        file_name="Store_Audit_Checklist_v1_dummy.xlsx",
        data_section="cash_deposit_pickups",
        column_headers=headers,
    )
    s.add(di_cd)
    s.flush()

    loaded_cd = 0
    for idx, row in enumerate(cd_rows, 1):
        sid = _resolve_sid(row, code_to_id, store_map) or first_store_id
        if not sid:
            continue
        s.add(CashDepositPickup(
            import_id=di_cd.id,
            store_id=sid,
            source_row_number=idx,
            sales_date=_parse_date(row.get("Sales Date (A)") or row.get("Sales Date")),
            cash_sales=_float(row.get("Cash Sales")),
            cash_deposited=_float(row.get("Cash Deposited (c)") or row.get("Cash Deposited")),
            difference=_float(row.get("Difference (B-C)") or row.get("Difference")),
            cms_pickup_date=_parse_date(row.get("CMS Pick Up date (D)") or row.get("CMS Pickup Date")),
            handover_delay_days=_int(row.get("Delay in cash handover (A-D)") or row.get("Handover Delay Days")),
            remarks=str(row.get("Remarks", "") or ""),
            raw_data=_json_safe(row),
        ))
        loaded_cd += 1
    print(f"  Loaded {loaded_cd} cash deposit pickup rows")

    # ====================================================================
    # 9. Expired Inventory from checklist sheet "3"
    # ====================================================================
    print("\n--- Loading expired inventory from checklist sheet '3' ---")
    headers, ei_rows = _read_named_sheet(CHECKLIST_FILE, "3")
    di_ei = DataImport(
        file_name="Store_Audit_Checklist_v1_dummy.xlsx",
        data_section="expired_inventory",
        column_headers=headers,
    )
    s.add(di_ei)
    s.flush()

    loaded_ei = 0
    for idx, row in enumerate(ei_rows, 1):
        sid = _resolve_sid(row, code_to_id, store_map) or first_store_id
        if not sid:
            continue
        s.add(ExpiredInventory(
            import_id=di_ei.id,
            store_id=sid,
            source_row_number=idx,
            article_code=str(row.get("Article Code", "") or ""),
            article_description=str(row.get("Article Description", "") or ""),
            expiry_date=_parse_date(row.get("Expiry Date")),
            review_date=_parse_date(row.get("Review Date")),
            quantity=_float(row.get("Quantity")),
            mrp=_float(row.get("MRP")),
            raw_data=_json_safe(row),
        ))
        loaded_ei += 1
    print(f"  Loaded {loaded_ei} expired inventory rows")

    # ====================================================================
    # 10. Create sample audits & issues so screens aren't empty
    # ====================================================================
    print("\n--- Creating sample audits ---")
    store_list = list(store_map.keys())
    audit_objs = {}
    random.seed(42)

    for i, sid in enumerate(store_list):
        aud_name = auditor_names[i % len(auditor_names)]
        auditor = auditors[aud_name]
        status = ["Completed", "Completed", "Ongoing", "Ongoing", "Planned",
                   "Planned", "Completed", "Ongoing", "Planned", "Ongoing"][i % 10]
        audit_id = f"AUD-{1000 + i}"
        scheduled = dt.datetime(2026, 6, random.randint(1, 28),
                                random.randint(6, 20), random.randint(0, 59))
        audit = Audit(
            id=audit_id,
            store_id=sid,
            checklist_id=full.id,
            auditor_id=auditor.id,
            created_by_id=admin.id,
            scheduled_at=scheduled,
            status=status,
            score=random.randint(65, 98) if status == "Completed" else None,
            submitted_at=scheduled if status == "Completed" else None,
        )
        s.add(audit)
        s.flush()
        audit_objs[audit_id] = audit

        # Plausible answers
        for q in questions.values():
            answer = None
            if status == "Completed":
                answer = random.choices(["Yes", "No", "Partial", "NA"],
                                        weights=[70, 10, 15, 5])[0]
            elif status == "Ongoing":
                answer = random.choice(["Yes", "Partial", None])
            s.add(AuditResponse(audit_id=audit.id, question_id=q.id, answer=answer))

    s.flush()
    print(f"  Created {len(audit_objs)} audits with responses")

    # ====================================================================
    # 11. Sample issues
    # ====================================================================
    print("\n--- Creating sample issues ---")
    issue_templates = [
        ("Freezer temperature fluctuation", "Critical", "Inventory Management"),
        ("Generator oil leakage detected", "Critical", "HR & Admin and Operations"),
        ("Fire safety certification pending", "Critical", "EHS & Compliance"),
        ("Wet floor in prep area", "High", "HR & Admin and Operations"),
        ("Delay in deposit of cash collected", "High", "Cashiering"),
        ("Excess cash discount provided", "Medium", "Cashiering"),
        ("Broken soap dispenser", "Low", "HR & Admin and Operations"),
        ("Expired stock not segregated", "High", "Inventory Management"),
        ("CCTV coverage gap", "High", "Inventory Management"),
        ("Staff ID cards not issued", "Low", "HR & Admin and Operations"),
        ("Ineligible SKUs returned", "Low", "Cashiering"),
    ]

    audit_list = list(audit_objs.values())
    for idx, (title, pri, proc) in enumerate(issue_templates):
        audit = audit_list[idx % len(audit_list)]
        store = store_map.get(audit.store_id)
        s.add(Issue(
            id=f"ISS{idx + 1:03d}",
            audit_id=audit.id,
            store_id=audit.store_id,
            title=title,
            description=f"Issue raised during audit {audit.id}",
            process=proc,
            priority=pri,
            status=random.choice(["Open", "In Progress", "Open", "Resolved"]),
            assignee_id=store.manager_id if store else None,
            raised_by_id=admin.id,
            due_date=(dt.date.today() + dt.timedelta(days=random.randint(-5, 14))),
            meta={"auto_raised": False},
        ))
    s.flush()
    print(f"  Created {len(issue_templates)} issues")

    # ====================================================================
    # 12. Auditor availability
    # ====================================================================
    auditor_list = list(auditors.values())
    if auditor_list:
        s.add(AuditorAvailability(
            auditor_id=auditor_list[0].id,
            from_date=dt.date.today() + dt.timedelta(days=3),
            to_date=dt.date.today() + dt.timedelta(days=6),
            reason="Planned leave",
        ))

    s.commit()

    # ---- summary -------------------------------------------------------
    print(f"\n{'='*60}")
    print(f"Seeded from Excel:")
    print(f"  {len(store_map)} stores, {len(questions)} questions")
    print(f"  {len(checklists) + 1} checklists, {len(audit_objs)} audits")
    print(f"  {loaded_scores} store scores, {loaded_obs} observations")
    print(f"  {loaded_cr} cash reconciliations")
    print(f"  {loaded_cd} cash deposit pickups")
    print(f"  {loaded_ei} expired inventory")
    print(f"  {len(issue_templates)} issues")
    print(f"  {len(auditors) + len(managers) + 1} users")
    print(f"\nLogins (password: {DEFAULT_PASSWORD})")
    print(f"  Audit Manager : am@retail-chain.com")
    for name, u in list(auditors.items())[:1]:
        print(f"  Auditor       : {u.email}")
    for name, u in list(managers.items())[:1]:
        print(f"  Store Manager : {u.email}")


if __name__ == "__main__":
    if "--reset" in sys.argv:
        Base.metadata.drop_all(bind=engine)
        print("Dropped all tables.")
    init_db()
    seed_from_excel()
